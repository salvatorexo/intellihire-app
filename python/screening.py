"""
IntelliHire - AI CV Screening Service (Render-ready, fast)
Uses pypdf (fast) instead of pdfplumber (slow).
Accepts cv_base64 (preferred) OR cv_path (for local testing).
"""

from flask import Flask, request, jsonify
from pypdf import PdfReader
import docx
import re
import os
import io
import base64
from collections import Counter
from math import log

app = Flask(__name__)

# ---------------- Text Extraction ----------------
def extract_pdf(path):
    """Fast PDF text extraction using pypdf."""
    text_parts = []
    reader = PdfReader(path)
    for page in reader.pages:
        try:
            t = page.extract_text() or ""
            if t:
                text_parts.append(t)
        except Exception:
            # Skip unreadable pages rather than crash
            continue
    return "\n".join(text_parts)

def extract_docx(path):
    doc = docx.Document(path)
    return "\n".join(p.text for p in doc.paragraphs)

def extract_text_from_bytes(raw_bytes, ext):
    """Extract text from in-memory bytes — faster than temp files."""
    ext = ext.lower()
    if ext == ".pdf":
        try:
            reader = PdfReader(io.BytesIO(raw_bytes))
            parts = []
            for page in reader.pages:
                try:
                    t = page.extract_text() or ""
                    if t:
                        parts.append(t)
                except Exception:
                    continue
            return "\n".join(parts)
        except Exception as e:
            return ""
    elif ext == ".docx":
        try:
            doc = docx.Document(io.BytesIO(raw_bytes))
            return "\n".join(p.text for p in doc.paragraphs)
        except Exception:
            return ""
    return ""

def extract_text_from_path(path):
    if path.lower().endswith(".pdf"):
        return extract_pdf(path)
    elif path.lower().endswith(".docx"):
        return extract_docx(path)
    return ""

# ---------------- Helpers ----------------
STOPWORDS = {
    "the","and","for","with","you","are","that","this","will","your","have",
    "from","our","job","work","all","any","can","but","not","was","were","has",
    "had","been","they","them","their","there","here","what","when","where",
    "which","who","how","why","about","into","over","under","also","more","most",
    "such","than","then","some","very","just","only","other","its","his","her",
    "who","whom","whose","been","being","does","did","doing","would","should",
    "could","may","might","must","shall","upon","within","without","during"
}

def tokenize(text):
    words = re.findall(r"\b[a-z][a-z0-9+#.]{2,}\b", text.lower())
    return [w for w in words if w not in STOPWORDS]

def term_frequencies(tokens):
    return Counter(tokens)

def tfidf_score(cv_tokens, job_tokens):
    cv_tf = term_frequencies(cv_tokens)
    job_tf = term_frequencies(job_tokens)

    vocab = set(cv_tf) | set(job_tf)
    if not vocab:
        return 0.0

    idf = {word: log(2 / (1 + (1 if word in cv_tf else 0) + (1 if word in job_tf else 0))) + 1
           for word in vocab}

    def vec(tf):
        return {w: (tf.get(w, 0) * idf[w]) for w in vocab}

    v1 = vec(cv_tf)
    v2 = vec(job_tf)

    dot = sum(v1[w] * v2[w] for w in vocab)
    mag1 = sum(v1[w] ** 2 for w in vocab) ** 0.5
    mag2 = sum(v2[w] ** 2 for w in vocab) ** 0.5

    if mag1 == 0 or mag2 == 0:
        return 0.0
    return dot / (mag1 * mag2)

def keyword_overlap(cv_tokens, job_tokens):
    cv_set = set(cv_tokens)
    job_set = set(job_tokens)
    if not job_set:
        return 0.0, []
    matched = cv_set & job_set
    return len(matched) / len(job_set), sorted(matched)[:25]

# ---------------- CV Text Loader ----------------
def load_cv_text(data):
    """
    Priority:
      1. cv_base64 + cv_filename -> decode & extract in-memory (fast)
      2. cv_path                 -> read from local disk (for local testing)
    Returns (text, error_message). On success, error_message is None.
    """
    if data.get("cv_base64"):
        try:
            raw = base64.b64decode(data["cv_base64"])
        except Exception as e:
            return "", f"base64 decode failed: {e}"

        filename = data.get("cv_filename") or "cv.pdf"
        ext = os.path.splitext(filename)[1].lower() or ".pdf"

        try:
            text = extract_text_from_bytes(raw, ext)
        except Exception as e:
            return "", f"CV extract failed: {e}"

        return text, None

    cv_path = data.get("cv_path")
    if cv_path:
        if not os.path.isfile(cv_path):
            return "", f"CV file not found at path: {cv_path}"
        try:
            return extract_text_from_path(cv_path), None
        except Exception as e:
            return "", f"CV extract failed: {e}"

    return "", "Missing cv_base64 or cv_path"

# ---------------- Endpoint ----------------
@app.route("/screen", methods=["POST"])
def screen():
    data = request.get_json(silent=True)
    if not data:
        return jsonify({"error": "Invalid JSON body"}), 400

    cv_text, err = load_cv_text(data)
    if err:
        return jsonify({
            "match_score": 0,
            "skills_score": 0,
            "education_score": 0,
            "experience_score": 0,
            "recommendation": "Weak Match",
            "matched_keywords": "",
            "note": err
        }), 200

    if not cv_text.strip():
        return jsonify({
            "match_score": 0,
            "skills_score": 0,
            "education_score": 0,
            "experience_score": 0,
            "recommendation": "Weak Match",
            "matched_keywords": "",
            "note": "CV text empty or unreadable"
        }), 200

    cv_tokens = tokenize(cv_text)

    job_text = " ".join([
        data.get("job_title", ""),
        data.get("job_description", ""),
        data.get("job_requirements", ""),
        data.get("job_skills", ""),
        data.get("education_requirements", ""),
    ])
    job_tokens = tokenize(job_text)

    tfidf = tfidf_score(cv_tokens, job_tokens)
    overlap_score, matched = keyword_overlap(cv_tokens, job_tokens)

    final_score = round(((tfidf * 0.6) + (overlap_score * 0.4)) * 100, 2)
    final_score = max(0.0, min(100.0, final_score))

    # ---- Sub-scores ----
    skills_tokens = tokenize(data.get("job_skills", ""))
    skills_score = keyword_overlap(cv_tokens, skills_tokens)[0] if skills_tokens else 0.0

    edu_tokens = tokenize(data.get("education_requirements", ""))
    edu_score = keyword_overlap(cv_tokens, edu_tokens)[0] if edu_tokens else 0.0

    exp_text = (data.get("job_description", "") + " " + data.get("job_requirements", ""))
    exp_tokens = tokenize(exp_text)
    exp_score = keyword_overlap(cv_tokens, exp_tokens)[0] if exp_tokens else 0.0

    # ---- Recommendation ----
    if final_score >= 80:
        recommendation = "Strong Match"
    elif final_score >= 60:
        recommendation = "Good Match"
    else:
        recommendation = "Weak Match"

    return jsonify({
        "match_score":      final_score,
        "skills_score":     round(skills_score * 100, 2),
        "education_score":  round(edu_score * 100, 2),
        "experience_score": round(exp_score * 100, 2),
        "recommendation":   recommendation,
        "matched_keywords": ", ".join(matched),
        "tfidf_score":      round(tfidf * 100, 2),
        "keyword_score":    round(overlap_score * 100, 2)
    })

@app.route("/", methods=["GET"])
def health():
    return jsonify({"status": "ok", "service": "IntelliHire AI Screening"})

# ---------------- Run ----------------
if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5001))
    print(f"🚀 IntelliHire AI Screening service running on port {port}")
    app.run(host="0.0.0.0", port=port, debug=False)
